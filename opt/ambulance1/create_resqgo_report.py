from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(r"C:\Users\lenovo\Desktop\ambulance 1")
PROJECT_ROOT = ROOT / "ambulance"
OUTPUT = ROOT / "ResQGo_Project_Report.docx"
IMAGE_PATH = PROJECT_ROOT / "Mobile_app" / "app" / "img.png"


def set_page_layout(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.left_margin = Inches(1.25)
    section.right_margin = Inches(1.0)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)


def set_default_style(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    fmt = normal.paragraph_format
    fmt.line_spacing = 1.5
    fmt.space_after = Pt(6)


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_end)


def add_footer_page_numbers(document: Document) -> None:
    for section in document.sections:
        footer = section.footer
        paragraph = footer.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_page_number(paragraph)


def add_title(document: Document, text: str, size: int = 18) -> None:
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)


def add_center_text(document: Document, text: str, size: int = 12, bold: bool = False) -> None:
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)


def add_chapter(document: Document, number: int, title: str) -> None:
    p1 = document.add_paragraph()
    p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r1 = p1.add_run(f"CHAPTER {number}")
    r1.bold = True
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(16)

    p2 = document.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(title)
    r2.bold = True
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(18)


def add_heading(document: Document, number: str, title: str, level: int = 1) -> None:
    size = 16 if level == 1 else 14
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(f"{number} {title}")
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)


def add_body(document: Document, text: str, justify: bool = True) -> None:
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if justify else WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)


def add_bullet(document: Document, text: str) -> None:
    p = document.add_paragraph(style="List Bullet")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)


def add_numbered(document: Document, text: str) -> None:
    p = document.add_paragraph(style="List Number")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)


def add_toc(document: Document) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = r'TOC \o "1-3" \h \z \u'
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Update field in Word to generate the table of contents."
    fld_sep.append(placeholder)
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(fld_end)


def add_table(document: Document, rows, widths=None):
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            cell = table.cell(i, j)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = value
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(11)
                    if i == 0:
                        run.bold = True
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = width
    document.add_paragraph()


def insert_cover_image(document: Document) -> None:
    if IMAGE_PATH.exists():
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(str(IMAGE_PATH), width=Inches(6.4))


def build_report() -> None:
    document = Document()
    set_page_layout(document)
    set_default_style(document)
    add_footer_page_numbers(document)

    add_center_text(document, "PROJECT REPORT", size=18, bold=True)
    add_center_text(document, "ON", size=14, bold=True)
    add_title(document, "ResQGo", size=22)
    add_center_text(document, "Smart Emergency Response and Ambulance Coordination System", size=15, bold=True)
    document.add_paragraph()
    insert_cover_image(document)
    document.add_paragraph()
    add_center_text(document, "Submitted in partial fulfillment of the academic project requirements", size=12)
    add_center_text(document, "Prepared from the implementation available in the ambulance 1 project workspace", size=12)
    document.add_paragraph()
    add_center_text(document, "Academic Year: 2025-2026", size=12, bold=True)
    document.add_page_break()

    add_center_text(document, "PROJECT REPORT", size=18, bold=True)
    add_title(document, "ResQGo", size=20)
    add_center_text(document, "Smart Emergency Response and Ambulance Coordination System", size=14, bold=True)
    add_body(
        document,
        "This report documents the design and implementation of ResQGo, a multi-component emergency response solution that combines an Android patient application, a Flask-based backend server, and a web dashboard for ambulance drivers. The report has been prepared in the structure requested by the supplied project format instructions.",
    )
    document.add_page_break()

    add_title(document, "ABSTRACT", size=18)
    add_body(
        document,
        "ResQGo is a smart emergency response system that enables patients to request ambulances with a single action, share real-time location, and track ambulance arrival. The solution combines a Jetpack Compose Android app, a Flask backend, and a Leaflet-based driver dashboard. It supports multilingual access, live GPS coordination, hospital discovery, and state-based mission updates. The system is designed to reduce response delay, improve communication between patients and drivers, and remain lightweight enough for practical deployment on common mobile devices.",
    )
    document.add_page_break()

    add_title(document, "ACKNOWLEDGEMENTS", size=18)
    add_body(
        document,
        "This report was prepared with the help of the source code, project notes, and formatting guidance available in the provided project folder and template documents. The implementation reflects the work put into building the Android patient experience, backend services, and web-based ambulance dashboard. The report also acknowledges the use of open technologies such as Android Jetpack Compose, Flask, Retrofit, and OpenStreetMap tools that made the project possible.",
    )
    document.add_page_break()

    add_title(document, "TABLE OF CONTENTS", size=18)
    add_toc(document)
    document.add_page_break()

    add_chapter(document, 1, "INTRODUCTION")
    add_heading(document, "1.1", "Project Overview", level=1)
    add_body(
        document,
        "ResQGo is a smart emergency response system designed to bridge the gap between patients and ambulance services. The project focuses on reducing response time through one-tap emergency requests, real-time GPS coordination, and clear communication between the patient side and the ambulance driver side. The system has been implemented as an Android patient application, a Flask-based backend service, and a browser-based ambulance dashboard.",
    )
    add_heading(document, "1.2", "Problem Statement", level=1)
    add_body(
        document,
        "During medical emergencies, delays are often caused by slow manual coordination, lack of location visibility, and difficulty in keeping patients informed after a request has been placed. Traditional ambulance booking processes also create barriers for users under stress. ResQGo addresses these issues by automating request submission, ambulance assignment visibility, and status tracking from request creation to hospital arrival.",
    )
    add_heading(document, "1.3", "Objectives", level=1)
    objectives = [
        "To provide a one-touch ambulance request mechanism for patients in urgent situations.",
        "To capture and share the user’s live location with the backend and ambulance driver.",
        "To support multilingual interaction for better accessibility in English, Hindi, and Kannada.",
        "To allow the patient to monitor ambulance status, distance, and estimated arrival time.",
        "To support hospital discovery and destination assignment once the patient is picked up.",
        "To build a lightweight and practical solution using open technologies and standard devices.",
    ]
    for item in objectives:
        add_bullet(document, item)
    add_heading(document, "1.4", "Scope of the Project", level=1)
    add_body(
        document,
        "The current project scope includes onboarding, patient session persistence, emergency request handling, driver-side mission management, ambulance location updates, live patient tracking, hospital search, and mission completion. The current implementation is suitable for academic demonstration and prototype deployment. Advanced traffic integration, analytics, and large-scale dispatch optimization are identified as future extensions.",
    )

    add_chapter(document, 2, "SYSTEM WORKFLOW AND FUNCTIONAL ANALYSIS")
    add_heading(document, "2.1", "User Workflow", level=1)
    add_heading(document, "2.1.1", "Phase 1: Onboarding and Authentication", level=2)
    for step in [
        "The application opens with a splash screen and checks whether the user session is already stored in shared preferences.",
        "New users choose a preferred language and continue through mobile-number entry and OTP verification.",
        "After verification, the user completes a profile using first name, last name, username, and password.",
        "Returning users can access the system through a login screen using phone details saved in the session flow.",
    ]:
        add_numbered(document, step)
    add_heading(document, "2.1.2", "Phase 2: Home Screen and Emergency Trigger", level=2)
    for step in [
        "The home screen requests fine-location permission and fetches the last known device location using the fused location provider.",
        "A large emergency button is visually emphasized so that the user can request help immediately during a crisis.",
        "A confirmation dialog is displayed to prevent accidental ambulance requests.",
    ]:
        add_numbered(document, step)
    add_heading(document, "2.1.3", "Phase 3: Backend Execution", level=2)
    for step in [
        "When the user confirms the request, the app sends a POST request containing patient identity, phone number, latitude, and longitude.",
        "The backend creates an emergency record, generates an emergency identifier, and stores the request state as pending.",
        "Available ambulance drivers can then view and accept the request through the web dashboard.",
    ]:
        add_numbered(document, step)
    add_heading(document, "2.1.4", "Phase 4: Live Status and Tracking", level=2)
    for step in [
        "The patient app polls the backend every three seconds for mission status changes such as pending, accepted, active, declined, or completed.",
        "Once accepted, the patient can view the assigned driver, ambulance number, driver phone number, travel distance, and ETA.",
        "A WebView-based Leaflet map is updated using the ambulance’s latest coordinates.",
        "After patient pickup, nearby hospitals are displayed and the destination becomes visible to the patient.",
    ]:
        add_numbered(document, step)
    add_heading(document, "2.2", "Key Functional Requirements", level=1)
    for item in [
        "Patient registration, login, and session persistence.",
        "Multilingual user interface rendering through a custom translation function.",
        "Location permission handling and GPS capture.",
        "Emergency request submission through REST APIs.",
        "Driver-side emergency polling, acceptance, and decline handling.",
        "Live ambulance movement updates and patient-facing map visualization.",
        "Hospital search and assignment during the active mission phase.",
        "Mission completion and state-based UI transitions.",
    ]:
        add_bullet(document, item)

    add_chapter(document, 3, "TECHNICAL ARCHITECTURE")
    add_heading(document, "3.1", "Overall Architecture", level=1)
    add_body(
        document,
        "The system follows a three-tier architecture. The Android mobile app serves as the patient-facing client. The Flask server exposes REST endpoints that coordinate emergency requests, ambulance location updates, hospital lookup, and mission state changes. The web dashboard serves ambulance drivers and consumes the same backend services. Database access is abstracted through a helper layer that attempts MySQL first and falls back to SQLite when necessary.",
    )
    add_heading(document, "3.2", "Technology Stack", level=1)
    add_table(
        document,
        [
            ["Layer", "Technology Used", "Purpose"],
            ["Mobile Client", "Kotlin, Jetpack Compose, Navigation Compose", "Patient onboarding, emergency request, live status UI"],
            ["Networking", "Retrofit with Gson Converter", "REST API communication from Android"],
            ["Location", "Google Play Services Fused Location Provider", "Current user location capture"],
            ["Backend", "Python Flask, Flask-SocketIO", "Business logic and API handling"],
            ["Database", "MySQL with SQLite fallback", "Persistent storage of drivers, emergencies, and location data"],
            ["Web Dashboard", "HTML, CSS, JavaScript, Leaflet", "Driver mission management and route display"],
            ["Maps and Geodata", "OpenStreetMap / Overpass API", "Hospital lookup and map visualization"],
        ],
    )
    add_heading(document, "3.3", "Android Application Design", level=1)
    add_body(
        document,
        "The Android application is implemented in a single Compose-driven activity with route-based navigation. Screens include Splash, Language Selection, Phone Entry, OTP Verification, Register, Login, Home, and Live Status. SharedPreferences is used for user session persistence. The translation function t() selects strings dynamically according to the chosen language, making the interface adaptable without maintaining separate resource bundles for each screen.",
    )
    add_heading(document, "3.4", "Backend Service Design", level=1)
    add_body(
        document,
        "The backend provides endpoints for driver login, driver registration, ambulance requests, emergency state retrieval, ambulance location retrieval, hospital discovery, hospital assignment, pickup confirmation, mission completion, system status, and logout. Emergency records move through a state model that includes pending, accepted, active, declined, and completed. The backend also computes approximate distance and ETA using the haversine formula based on ambulance and target coordinates.",
    )
    add_heading(document, "3.5", "Web Dashboard Design", level=1)
    add_body(
        document,
        "The driver dashboard is implemented in JavaScript and uses Leaflet to present the map, patient location, hospital destination, and route path. It polls for new emergencies, supports accept and decline actions, performs route calculation, sends location updates to the backend, and updates mission details on screen. The dashboard supports both the pickup phase and the hospital transport phase.",
    )

    add_chapter(document, 4, "MODULE DESCRIPTION")
    add_heading(document, "4.1", "Authentication and Session Module", level=1)
    add_body(
        document,
        "This module manages splash-screen checking, user language selection, OTP-based verification flow, registration, and session persistence. SharedPreferences stores user profile details locally so that the app can reopen directly into the home screen when a valid session exists.",
    )
    add_heading(document, "4.2", "Emergency Request Module", level=1)
    add_body(
        document,
        "This module gathers the patient’s current location and sends an ambulance request payload to the backend. The payload includes the user’s name, phone number, and precise latitude and longitude. A confirmation dialog is included to reduce accidental usage during normal app interaction.",
    )
    add_heading(document, "4.3", "Live Tracking Module", level=1)
    add_body(
        document,
        "The live tracking module repeatedly queries the backend for state changes and current ambulance coordinates. On the patient side, the map is rendered through an embedded WebView running a lightweight Leaflet map. On the driver side, route coordinates are used to animate ambulance movement and update the mission path toward the patient and later toward the hospital.",
    )
    add_heading(document, "4.4", "Hospital Search Module", level=1)
    add_body(
        document,
        "Once a patient is picked up, the driver can search nearby hospitals. The backend calls the Overpass API to locate hospitals around the selected coordinates and returns structured data to the dashboard. If the external service is unavailable, the backend returns fallback hospital entries so the workflow remains functional during demonstration.",
    )
    add_heading(document, "4.5", "Database and Persistence Module", level=1)
    add_body(
        document,
        "The persistence module stores drivers, emergencies, and ambulance location updates. The database helper tries a MySQL connection first and automatically switches to SQLite if MySQL is unavailable. This design improves portability and allows the application to run in classroom or lab setups without demanding a fixed database server.",
    )

    add_chapter(document, 5, "INTERFACE AND DATA DESIGN")
    add_heading(document, "5.1", "User Interface Highlights", level=1)
    for item in [
        "The patient application uses a high-contrast, emergency-oriented visual style with a prominent red panic button.",
        "The language-aware screens improve usability for regional users by supporting English, Hindi, and Kannada.",
        "The live status screen displays driver details, ETA, distance, and hospital list in a consolidated layout.",
        "The web dashboard provides a route map, emergency popup, mission status panel, and hospital selection controls.",
    ]:
        add_bullet(document, item)
    add_heading(document, "5.2", "Key API Endpoints", level=1)
    add_table(
        document,
        [
            ["Endpoint", "Method", "Function"],
            ["/api/request_ambulance", "POST", "Creates a new emergency request"],
            ["/api/emergency_status", "GET", "Returns current emergency state and driver details"],
            ["/api/ambulance_location", "GET", "Returns live ambulance coordinates, distance, and ETA"],
            ["/api/nearby_hospitals", "GET", "Returns nearby hospitals around a coordinate"],
            ["/api/accept_emergency", "POST", "Marks an emergency as accepted by a driver"],
            ["/api/send_location", "POST", "Stores a new ambulance location update"],
            ["/api/assign_hospital", "POST", "Assigns a hospital destination to the mission"],
            ["/api/complete_mission", "POST", "Marks the current mission as completed"],
        ],
    )
    add_heading(document, "5.3", "Representative Data Entities", level=1)
    add_table(
        document,
        [
            ["Entity", "Important Fields", "Purpose"],
            ["drivers", "driver_name, username, password, ambulance_no, phone", "Stores ambulance driver credentials and profile data"],
            ["emergencies", "patient_name, patient_mobile, lat, lon, status, driver_id, otp, dest_name", "Stores ambulance request and mission state"],
            ["ambulance_locations", "driver_id, lat, lon, timestamp", "Stores location updates from driver dashboard"],
        ],
    )

    add_chapter(document, 6, "IMPLEMENTATION SUMMARY")
    add_heading(document, "6.1", "Important Implementation Decisions", level=1)
    for item in [
        "Jetpack Compose was used for declarative UI and quick screen iteration on Android.",
        "Retrofit and Gson were selected for concise mobile networking code.",
        "Leaflet inside a WebView was chosen to keep patient-side tracking lightweight without a heavy native map SDK.",
        "Backend fallback from MySQL to SQLite improves resilience during local development and demonstrations.",
        "Polling is used for state updates, which keeps the architecture simple and easy to test.",
    ]:
        add_bullet(document, item)
    add_heading(document, "6.2", "Innovation Highlights", level=1)
    for item in [
        "Zero-latency interaction path from launch to emergency request.",
        "Multilingual accessibility for emergency use cases.",
        "Low-resource map rendering suitable for entry-level devices.",
        "Graceful handling of state transitions such as decline, pickup, hospital assignment, and completion.",
        "Fallback hospital data to preserve workflow continuity when external map services fail.",
    ]:
        add_bullet(document, item)

    add_chapter(document, 7, "TESTING, RESULTS, AND LIMITATIONS")
    add_heading(document, "7.1", "Observed Functional Behavior", level=1)
    add_body(
        document,
        "Source inspection shows that the prototype covers the end-to-end emergency flow: session management, emergency request creation, driver acceptance, live location updates, patient tracking, nearby hospital display, and mission completion. The Android application and the driver dashboard are connected through REST interfaces exposed by the Flask backend.",
    )
    add_heading(document, "7.2", "Strengths of the Current Prototype", level=1)
    for item in [
        "Clear separation between patient, backend, and driver responsibilities.",
        "Readable and demonstrable workflow for academic presentation.",
        "Lightweight open-source map integration.",
        "Practical fallback design for data storage and hospital results.",
    ]:
        add_bullet(document, item)
    add_heading(document, "7.3", "Current Limitations", level=1)
    for item in [
        "The patient OTP verification flow in the current app is UI-based and not fully validated against the backend.",
        "The Android API base URL is configured for a specific local network IP and may need adjustment in other environments.",
        "The backend uses plain credential matching in the shown implementation and would benefit from stronger password security for production use.",
        "Polling is simple but less efficient than event-driven live synchronization for large-scale deployment.",
    ]:
        add_bullet(document, item)

    add_chapter(document, 8, "CONCLUSION AND FUTURE SCOPE")
    add_heading(document, "8.1", "Conclusion", level=1)
    add_body(
        document,
        "ResQGo demonstrates how a practical emergency response workflow can be built by combining a mobile client, backend coordination service, and web dashboard. The project successfully addresses the major academic goals of ambulance request automation, real-time tracking, hospital visibility, and multilingual accessibility. Its implementation choices show a strong focus on usability, low deployment cost, and clear operational flow.",
    )
    add_heading(document, "8.2", "Future Scope", level=1)
    for item in [
        "Secure authentication with encrypted passwords and token-based sessions.",
        "Automatic nearest-ambulance selection using availability and route intelligence.",
        "Push-based live updates using sockets for both patient and driver.",
        "Integration with hospitals, emergency contacts, and traffic control systems.",
        "Analytics dashboard for emergency trends and response-time measurement.",
    ]:
        add_bullet(document, item)

    add_chapter(document, 9, "REFERENCES")
    references = [
        "[1] Android Developers, Jetpack Compose Documentation, Google Developers.",
        "[2] Square, Retrofit Documentation, REST Client for Android and Java.",
        "[3] Flask Documentation, Pallets Projects.",
        "[4] OpenStreetMap Foundation, OpenStreetMap and Overpass API Documentation.",
        "[5] Leaflet Documentation, JavaScript Library for Interactive Maps.",
        "[6] Project source files from the ambulance 1 workspace: Android client, Flask backend, and driver dashboard implementation.",
    ]
    for ref in references:
        add_body(document, ref, justify=False)

    document.save(str(OUTPUT))


if __name__ == "__main__":
    build_report()
